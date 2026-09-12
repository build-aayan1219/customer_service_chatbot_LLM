import hashlib
import json
import logging
import re
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from langchain_core.documents import Document
from langchain_community.vectorstores import FAISS

from src.langchain_helper import get_embeddings


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

CONVERSATIONS_DIR = (
    BASE_DIR
    / "knowledge_base"
    / "conversations"
)


# ============================================================
# CONFIGURATION
# ============================================================

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "txt",
    "csv",
}

MAX_FILE_SIZE_MB = 25
MAX_FILES_PER_CHAT = 10

CHUNK_SIZE = 1800
CHUNK_OVERLAP = 250

MAX_RETRIEVAL_CANDIDATES = 8


# ============================================================
# LOGGING
# ============================================================

logger = logging.getLogger(__name__)


# ============================================================
# DIRECTORY HELPERS
# ============================================================

def get_chat_directory(chat_id):
    return (
        CONVERSATIONS_DIR
        / str(chat_id)
    )


def get_files_directory(chat_id):
    return (
        get_chat_directory(chat_id)
        / "files"
    )


def get_index_directory(chat_id):
    return (
        get_chat_directory(chat_id)
        / "index"
    )


def get_manifest_path(chat_id):
    return (
        get_chat_directory(chat_id)
        / "manifest.json"
    )


def ensure_chat_directories(chat_id):
    get_files_directory(chat_id).mkdir(
        parents=True,
        exist_ok=True,
    )

    get_index_directory(chat_id).mkdir(
        parents=True,
        exist_ok=True,
    )


# ============================================================
# MANIFEST
# ============================================================

def load_manifest(chat_id):
    ensure_chat_directories(chat_id)

    manifest_path = get_manifest_path(chat_id)

    if not manifest_path.exists():
        return {}

    try:
        with manifest_path.open(
            "r",
            encoding="utf-8",
        ) as file:
            data = json.load(file)

        if isinstance(data, dict):
            return data

    except (
        json.JSONDecodeError,
        OSError,
    ) as error:

        logger.warning(
            "Could not load conversation manifest: %s",
            error,
        )

    return {}


def save_manifest(chat_id, manifest):
    ensure_chat_directories(chat_id)

    manifest_path = get_manifest_path(chat_id)

    with manifest_path.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            manifest,
            file,
            indent=2,
            ensure_ascii=False,
        )


# ============================================================
# FILE HASH
# ============================================================

def calculate_file_hash(file_data):
    return hashlib.sha256(
        file_data
    ).hexdigest()


# ============================================================
# SAFE FILE NAME
# ============================================================

def sanitize_filename(filename):
    filename = Path(
        str(filename)
    ).name

    filename = re.sub(
        r"[^A-Za-z0-9._() -]",
        "_",
        filename,
    )

    filename = filename.strip()

    if not filename:
        filename = "uploaded_file"

    return filename


# ============================================================
# FILE SIZE
# ============================================================

def format_file_size(size_bytes):
    if size_bytes < 1024:
        return f"{size_bytes} B"

    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"

    return f"{size_bytes / (1024 * 1024):.1f} MB"


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf(file_path):
    from pypdf import PdfReader

    reader = PdfReader(
        str(file_path)
    )

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        text = (
            page.extract_text()
            or ""
        ).strip()

        if text:
            pages.append(
                (
                    page_number,
                    text,
                )
            )

    return pages


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx(file_path):
    from docx import Document as DocxDocument

    document = DocxDocument(
        str(file_path)
    )

    sections = []
    current_section = []
    current_heading = None

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        style_name = ""

        try:
            style_name = (
                paragraph.style.name
                or ""
            ).lower()
        except Exception:
            pass

        is_heading = (
            "heading" in style_name
        )

        if is_heading:

            if current_section:

                sections.append(
                    (
                        current_heading,
                        "\n".join(
                            current_section
                        ),
                    )
                )

            current_heading = text
            current_section = []

        else:

            current_section.append(
                text
            )

    if current_section:

        sections.append(
            (
                current_heading,
                "\n".join(
                    current_section
                ),
            )
        )

    if not sections:
        return []

    results = []

    for heading, text in sections:

        if heading:

            combined = (
                f"{heading}\n{text}"
            )

        else:

            combined = text

        results.append(
            (
                1,
                combined.strip(),
            )
        )

    return results


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt(file_path):
    text = file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    ).strip()

    if not text:
        return []

    return [
        (
            1,
            text,
        )
    ]


# ============================================================
# CSV EXTRACTION
# ============================================================

def extract_csv(file_path):
    dataframe = pd.read_csv(
        file_path
    )

    rows = []

    dataframe = dataframe.fillna("")

    for row_number, row in dataframe.iterrows():

        parts = []

        for column in dataframe.columns:

            value = str(
                row[column]
            ).strip()

            if value:

                parts.append(
                    f"{column}: {value}"
                )

        text = "\n".join(
            parts
        ).strip()

        if text:

            rows.append(
                (
                    row_number + 2,
                    text,
                )
            )

    return rows


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_file(file_path):

    extension = (
        file_path
        .suffix
        .lower()
        .lstrip(".")
    )

    if extension == "pdf":
        return extract_pdf(
            file_path
        )

    if extension == "docx":
        return extract_docx(
            file_path
        )

    if extension == "txt":
        return extract_txt(
            file_path
        )

    if extension == "csv":
        return extract_csv(
            file_path
        )

    raise ValueError(
        f"Unsupported file type: .{extension}"
    )


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):
    text = str(
        text or ""
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


# ============================================================
# HEADING DETECTION
# ============================================================

def is_probable_heading(line):
    line = line.strip()

    if not line:
        return False

    if len(line) > 120:
        return False

    if re.match(
        r"^\d+[\.\)]\s+",
        line,
    ):
        return True

    if re.match(
        r"^[A-Z][A-Za-z0-9\s\-:&/]{2,80}:?$",
        line,
    ) and len(
        line.split()
    ) <= 12:

        return True

    return False


# ============================================================
# CHUNKING
# ============================================================

def split_text(
    text,
    chunk_size=CHUNK_SIZE,
    overlap=CHUNK_OVERLAP,
):
    text = str(
        text or ""
    ).strip()

    if not text:
        return []

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if not lines:
        return []

    sections = []

    current_heading = None
    current_lines = []

    for line in lines:

        if is_probable_heading(line):

            if current_lines:

                sections.append(
                    (
                        current_heading,
                        "\n".join(
                            current_lines
                        ).strip(),
                    )
                )

            current_heading = line
            current_lines = []

        else:

            current_lines.append(
                line
            )

    if current_lines:

        sections.append(
            (
                current_heading,
                "\n".join(
                    current_lines
                ).strip(),
            )
        )

    if not sections:

        sections = [
            (
                None,
                text,
            )
        ]

    chunks = []

    for heading, section_text in sections:

        if not section_text:
            continue

        if heading:

            section_text = (
                f"{heading}\n"
                f"{section_text}"
            )

        if len(section_text) <= chunk_size:

            chunks.append(
                section_text
            )

            continue

        words = section_text.split()

        start = 0

        while start < len(words):

            current_words = []

            current_length = 0

            index = start

            while index < len(words):

                word = words[index]

                additional_length = (
                    len(word)
                    + (
                        1
                        if current_words
                        else 0
                    )
                )

                if (
                    current_words
                    and
                    current_length
                    + additional_length
                    > chunk_size
                ):

                    break

                current_words.append(
                    word
                )

                current_length += (
                    additional_length
                )

                index += 1

            chunk = " ".join(
                current_words
            ).strip()

            if chunk:
                chunks.append(
                    chunk
                )

            if index >= len(words):
                break

            overlap_words = max(
                1,
                int(
                    overlap
                    / max(
                        len(words[index - 1]),
                        1,
                    )
                ),
            )

            start = max(
                index - overlap_words,
                start + 1,
            )

    return [
        chunk.strip()
        for chunk in chunks
        if chunk.strip()
    ]


# ============================================================
# LANGCHAIN DOCUMENT CREATION
# ============================================================

def create_documents(
    file_path,
    chat_id,
):
    extracted = extract_file(
        file_path
    )

    documents = []

    file_name = file_path.name

    chunk_counter = 0

    for location, text in extracted:

        chunks = split_text(
            text
        )

        for chunk in chunks:

            chunk_counter += 1

            lines = [
                line.strip()
                for line
                in chunk.splitlines()
                if line.strip()
            ]

            section = ""

            if lines:

                first_line = lines[0]

                if (
                    len(first_line) <= 120
                    and (
                        is_probable_heading(
                            first_line
                        )
                        or re.match(
                            r"^\d+[\.\)]",
                            first_line,
                        )
                    )
                ):

                    section = first_line

            metadata = {
                "source_file": file_name,
                "file_name": file_name,
                "filename": file_name,
                "source_type": "conversation_file",
                "chat_id": str(chat_id),
                "location": (
                    f"Page/row {location}, "
                    f"chunk {chunk_counter}"
                ),
                "section": section,
                "chunk": chunk_counter,
            }

            documents.append(
                Document(
                    page_content=chunk,
                    metadata=metadata,
                )
            )

    return documents


# ============================================================
# LOAD CONVERSATION VECTOR DB
# ============================================================

def load_conversation_vector_db(chat_id):
    index_directory = (
        get_index_directory(
            chat_id
        )
    )

    index_file = (
        index_directory
        / "index.faiss"
    )

    pickle_file = (
        index_directory
        / "index.pkl"
    )

    if (
        not index_file.exists()
        or not pickle_file.exists()
    ):

        return None

    try:

        vectordb = FAISS.load_local(
            str(index_directory),
            get_embeddings(),
            allow_dangerous_deserialization=True,
        )

        return vectordb

    except Exception as error:

        logger.error(
            "Could not load conversation FAISS for %s: %s",
            chat_id,
            error,
        )

        return None


# ============================================================
# SAVE / UPDATE CONVERSATION VECTOR DB
# ============================================================

def add_documents_to_index(
    chat_id,
    documents,
):
    if not documents:
        return None

    ensure_chat_directories(
        chat_id
    )

    index_directory = (
        get_index_directory(
            chat_id
        )
    )

    existing_db = (
        load_conversation_vector_db(
            chat_id
        )
    )

    embeddings = get_embeddings()

    if existing_db is None:

        vectordb = FAISS.from_documents(
            documents,
            embeddings,
        )

    else:

        existing_db.add_documents(
            documents
        )

        vectordb = existing_db

    vectordb.save_local(
        str(index_directory)
    )

    return vectordb


# ============================================================
# ATTACHMENT SIGNATURE
# ============================================================

def get_attachment_signature(
    uploaded_files
):
    if not uploaded_files:
        return ""

    parts = []

    for uploaded_file in uploaded_files:

        try:

            data = uploaded_file.getvalue()

            file_hash = (
                calculate_file_hash(
                    data
                )
            )

            parts.append(
                f"{uploaded_file.name}:"
                f"{len(data)}:"
                f"{file_hash}"
            )

        except Exception:

            parts.append(
                str(
                    uploaded_file.name
                )
            )

    return "|".join(
        sorted(parts)
    )


# ============================================================
# LIST CONVERSATION FILES
# ============================================================

def list_conversation_files(
    chat_id
):
    manifest = load_manifest(
        chat_id
    )

    records = []

    for file_id, record in manifest.items():

        if not isinstance(
            record,
            dict,
        ):
            continue

        records.append(
            {
                "id": file_id,
                **record,
            }
        )

    records.sort(
        key=lambda item:
        item.get(
            "uploaded_at",
            "",
        )
    )

    return records


# ============================================================
# ADD CONVERSATION FILES
# ============================================================

def add_conversation_files(
    chat_id,
    uploaded_files,
):
    ensure_chat_directories(
        chat_id
    )

    manifest = load_manifest(
        chat_id
    )

    result = {
        "added": [],
        "skipped": [],
        "errors": [],
    }

    if not uploaded_files:
        return result

    existing_count = len(
        manifest
    )

    for uploaded_file in uploaded_files:

        file_name = sanitize_filename(
            uploaded_file.name
        )

        extension = (
            Path(file_name)
            .suffix
            .lower()
            .lstrip(".")
        )

        if extension not in ALLOWED_EXTENSIONS:

            result["errors"].append(
                f"{file_name}: "
                f"Unsupported file type."
            )

            continue

        try:

            file_data = (
                uploaded_file.getvalue()
            )

        except Exception as error:

            result["errors"].append(
                f"{file_name}: "
                f"Could not read file ({error})."
            )

            continue

        size_bytes = len(
            file_data
        )

        if (
            size_bytes
            > MAX_FILE_SIZE_MB
            * 1024
            * 1024
        ):

            result["errors"].append(
                f"{file_name}: "
                f"File exceeds the "
                f"{MAX_FILE_SIZE_MB} MB limit."
            )

            continue

        file_hash = (
            calculate_file_hash(
                file_data
            )
        )

        duplicate = False

        for record in manifest.values():

            if (
                record.get(
                    "file_hash"
                )
                == file_hash
            ):

                duplicate = True
                break

        if duplicate:

            result["skipped"].append(
                {
                    "name": file_name,
                    "reason": "duplicate",
                }
            )

            continue

        if (
            existing_count
            >= MAX_FILES_PER_CHAT
        ):

            result["errors"].append(
                f"Maximum of "
                f"{MAX_FILES_PER_CHAT} "
                f"files per conversation "
                f"is allowed."
            )

            break

        file_id = file_hash[:16]

        file_path = (
            get_files_directory(
                chat_id
            )
            / file_name
        )

        try:

            with file_path.open(
                "wb"
            ) as file:

                file.write(
                    file_data
                )

            documents = create_documents(
                file_path,
                chat_id,
            )

            if not documents:

                file_path.unlink(
                    missing_ok=True
                )

                result["errors"].append(
                    f"{file_name}: "
                    f"No readable text was found."
                )

                continue

            add_documents_to_index(
                chat_id,
                documents,
            )

            uploaded_at = (
                datetime.utcnow()
                .isoformat()
                + "Z"
            )

            record = {
                "file_name": file_name,
                "original_name": uploaded_file.name,
                "file_hash": file_hash,
                "size_bytes": size_bytes,
                "size_display": format_file_size(
                    size_bytes
                ),
                "chunk_count": len(
                    documents
                ),
                "uploaded_at": uploaded_at,
                "extension": extension,
            }

            manifest[file_id] = record

            existing_count += 1

            result["added"].append(
                {
                    "id": file_id,
                    **record,
                }
            )

        except Exception as error:

            logger.exception(
                "Failed processing conversation file %s",
                file_name,
            )

            try:

                file_path.unlink(
                    missing_ok=True
                )

            except Exception:
                pass

            result["errors"].append(
                f"{file_name}: "
                f"Processing failed: {error}"
            )

    save_manifest(
        chat_id,
        manifest,
    )

    return result


# ============================================================
# REMOVE SINGLE CONVERSATION FILE
# ============================================================

def remove_conversation_file(
    chat_id,
    file_id,
):
    manifest = load_manifest(
        chat_id
    )

    record = manifest.get(
        str(file_id)
    )

    if not record:
        return False

    file_name = record.get(
        "file_name"
    )

    if file_name:

        file_path = (
            get_files_directory(
                chat_id
            )
            / file_name
        )

        try:

            file_path.unlink(
                missing_ok=True
            )

        except Exception as error:

            logger.warning(
                "Could not remove file %s: %s",
                file_path,
                error,
            )

    del manifest[
        str(file_id)
    ]

    save_manifest(
        chat_id,
        manifest,
    )

    rebuild_conversation_index(
        chat_id
    )

    return True


# ============================================================
# REBUILD CONVERSATION INDEX
# ============================================================

def rebuild_conversation_index(
    chat_id
):
    ensure_chat_directories(
        chat_id
    )

    manifest = load_manifest(
        chat_id
    )

    all_documents = []

    files_directory = (
        get_files_directory(
            chat_id
        )
    )

    for record in manifest.values():

        file_name = record.get(
            "file_name"
        )

        if not file_name:
            continue

        file_path = (
            files_directory
            / file_name
        )

        if not file_path.exists():
            continue

        try:

            documents = create_documents(
                file_path,
                chat_id,
            )

            all_documents.extend(
                documents
            )

        except Exception as error:

            logger.warning(
                "Could not rebuild %s: %s",
                file_name,
                error,
            )

    index_directory = (
        get_index_directory(
            chat_id
        )
    )

    if not all_documents:

        if index_directory.exists():

            shutil.rmtree(
                index_directory,
                ignore_errors=True,
            )

        index_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        return None

    vectordb = FAISS.from_documents(
        all_documents,
        get_embeddings(),
    )

    vectordb.save_local(
        str(index_directory)
    )

    return vectordb


# ============================================================
# SEARCH CONVERSATION FILES
# ============================================================

def search_conversation_files(
    chat_id,
    query,
    k=MAX_RETRIEVAL_CANDIDATES,
):
    vectordb = (
        load_conversation_vector_db(
            chat_id
        )
    )

    if vectordb is None:
        return []

    query = str(
        query or ""
    ).strip()

    if not query:
        return []

    try:

        results = (
            vectordb
            .similarity_search_with_relevance_scores(
                query,
                k=k,
            )
        )

        return results

    except Exception as error:

        logger.warning(
            "Conversation-file similarity search failed: %s",
            error,
        )

        try:

            documents = (
                vectordb.similarity_search(
                    query,
                    k=k,
                )
            )

            return [
                (
                    document,
                    0.5,
                )
                for document in documents
            ]

        except Exception as fallback_error:

            logger.error(
                "Conversation-file search failed: %s",
                fallback_error,
            )

            return []



# ============================================================
# GET ALL CONVERSATION DOCUMENTS
# ============================================================

def get_conversation_documents(
    chat_id,
):
    vectordb = load_conversation_vector_db(chat_id)

    if vectordb is None:
        return []

    try:
        documents = list(vectordb.docstore._dict.values())
    except Exception as error:
        logger.warning(
            "Could not read conversation documents: %s",
            error,
        )
        return []

    return [
        document
        for document in documents
        if isinstance(document, Document)
    ]


# ============================================================
# FILE OVERVIEW DOCUMENTS
# ============================================================

def get_conversation_overview_documents(
    chat_id,
    max_chunks=10,
):
    """Return representative chunks from all files attached to a chat."""
    documents = get_conversation_documents(chat_id)

    if not documents:
        return []

    grouped = {}
    for document in documents:
        metadata = document.metadata if isinstance(document.metadata, dict) else {}
        file_name = (
            metadata.get("source_file")
            or metadata.get("file_name")
            or metadata.get("filename")
            or "unknown"
        )
        grouped.setdefault(str(file_name), []).append(document)

    selected = []
    number_of_files = max(len(grouped), 1)
    chunks_per_file = max(2, max_chunks // number_of_files)

    for file_name, file_documents in grouped.items():
        file_documents.sort(
            key=lambda document: int(
                document.metadata.get(
                    "chunk_index",
                    document.metadata.get("chunk", 0),
                ) or 0
            )
        )

        count = len(file_documents)
        if count <= chunks_per_file:
            chosen = file_documents
        else:
            indexes = [0]

            # Prefer chunks containing headings/section labels because they
            # give the model more semantic coverage for broad questions.
            heading_indexes = []
            for index, document in enumerate(file_documents):
                metadata = document.metadata if isinstance(document.metadata, dict) else {}
                section = str(metadata.get("section") or "").strip()
                if section:
                    heading_indexes.append(index)

            for index in heading_indexes:
                if len(indexes) >= chunks_per_file:
                    break
                if index not in indexes:
                    indexes.append(index)

            remaining = chunks_per_file - len(indexes)
            if remaining > 0:
                for i in range(1, remaining + 1):
                    position = round(
                        i * (count - 1) / max(remaining, 1)
                    )
                    if position not in indexes:
                        indexes.append(position)

            chosen = [file_documents[index] for index in sorted(set(indexes))]

        total = len(file_documents)
        for document in chosen:
            metadata = document.metadata if isinstance(document.metadata, dict) else {}
            metadata.setdefault("source_file", file_name)
            metadata.setdefault("file_name", file_name)
            metadata.setdefault("filename", file_name)
            metadata["source_type"] = "conversation_file"
            metadata["retrieval_source"] = "conversation_file"
            metadata["overview_retrieval"] = True
            metadata["total_chunks"] = total
            document.metadata = metadata
            selected.append(document)

    return selected[:max_chunks]

# ============================================================
# DELETE ALL FILES FOR ONE CHAT
# ============================================================

def delete_conversation_files(
    chat_id
):
    chat_directory = (
        get_chat_directory(
            chat_id
        )
    )

    if chat_directory.exists():

        shutil.rmtree(
            chat_directory,
            ignore_errors=True,
        )


# ============================================================
# DELETE ALL CONVERSATION FILES
# ============================================================

def delete_all_conversation_files(
    chat_ids=None
):
    if chat_ids is None:

        if CONVERSATIONS_DIR.exists():

            shutil.rmtree(
                CONVERSATIONS_DIR,
                ignore_errors=True,
            )

        return

    for chat_id in chat_ids:

        delete_conversation_files(
            chat_id
        )